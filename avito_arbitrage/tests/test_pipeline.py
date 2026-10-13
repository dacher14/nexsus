import json

from openpyxl import load_workbook

from arbitrage.dedup import deduplicate
from arbitrage.export.workbook import ALL_DEALS_COLUMNS
from arbitrage.models import Listing, ListingStatus, listing_id_from_url
from arbitrage.pipeline import run_campaign
from arbitrage.sources import AvitoAuthorizedSource, SourceNotConfigured, load_listings
from conftest import make_listing

import pytest


def test_listing_id_from_avito_url():
    assert listing_id_from_url("https://www.avito.ru/moskva/telefony/iphone_12_128gb_1234567890") == "1234567890"
    assert listing_id_from_url("https://www.avito.ru/moskva/telefony/iphone_12_128gb_1234567890?context=x") == "1234567890"
    assert listing_id_from_url("https://example.com/x").startswith("u")


def test_price_parsing():
    l = Listing.from_dict({"url": "https://example.invalid/a_1000001", "title": "X", "price": "12 500 ₽"})
    assert l.price == 12500


def test_dedup_exact_and_reposts():
    a = make_listing("iPhone 12 128GB", 15000, seller_id="s1")
    a_again = Listing.from_dict({**a.to_dict(), "search_query": "второй запрос"})
    repost = make_listing("iPhone 12 128GB", 15000, seller_id="s1")
    shop_other_unit = make_listing("iPhone 12 128GB", 17000, seller_id="s1")
    other = make_listing("iPhone 12 128GB", 15000, seller_id="s2")
    res = deduplicate([a, a_again, repost, shop_other_unit, other])
    assert res.exact_duplicates == 1
    assert [r[0] for r in res.reposts] == [repost.listing_id]
    assert len(res.unique) == 3


def test_unconfigured_avito_source_is_explicit():
    with pytest.raises(SourceNotConfigured):
        list(AvitoAuthorizedSource().fetch())


def test_invalid_records_are_rejected(tmp_path):
    p = tmp_path / "x.json"
    p.write_text(json.dumps([{"title": "без ссылки", "price": 100},
                             {"url": "https://example.invalid/a_1234567", "title": "ok", "price": 500}]),
                 encoding="utf-8")
    from arbitrage.sources import JsonFileSource
    listings, errors = load_listings([JsonFileSource(p)])
    assert len(listings) == 1 and len(errors) == 1


def test_campaign_end_to_end(storage, settings, synthetic_file, tmp_path):
    out = tmp_path / "campaign.xlsx"
    result = run_campaign(storage, settings, [synthetic_file], out_path=out)
    assert result.invalid == 2
    assert result.exact_duplicates == 15
    assert 200 <= result.unique <= 400
    assert len(result.analyses) == result.unique
    assert any(a.critical for a in result.analyses)

    wb = load_workbook(out)
    assert wb.sheetnames == ["ALL DEALS", "TOP OPPORTUNITIES", "REPAIR LAB", "SELLER NEGOTIATIONS",
                             "MARKET INTELLIGENCE", "SETTINGS & CAPITAL"]
    ws = wb["ALL DEALS"]
    headers = [c.value for c in ws[1]]
    assert headers == [c.header for c in ALL_DEALS_COLUMNS] and len(headers) == 54
    assert ws.max_row == result.unique + 1
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref
    total_col = headers.index("Полная себестоимость") + 1
    formulas = [ws.cell(row=r, column=total_col).value for r in range(2, 20)]
    assert any(isinstance(f, str) and f.startswith("=J") for f in formulas)
    validations = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert any("APPROVED" in f for f in validations)
    assert "ТЕСТОВЫЕ ДАННЫЕ" in str(wb["SETTINGS & CAPITAL"]["A1"].value)
    top = wb["TOP OPPORTUNITIES"]
    assert 2 < top.max_row <= settings.top_size + 2


def test_recheck_tracks_price_history_and_unavailable(storage, settings, tmp_path):
    first = [make_listing("iPhone 12 128GB", p, "Всё работает") for p in (15000, 16000, 17000)]
    path = tmp_path / "a.json"
    path.write_text(json.dumps([l.to_dict() for l in first], ensure_ascii=False), encoding="utf-8")
    run_campaign(storage, settings, [path])

    changed = first[0].to_dict() | {"price": 14000}
    path.write_text(json.dumps([changed, first[1].to_dict()], ensure_ascii=False), encoding="utf-8")
    result = run_campaign(storage, settings, [path], full_recheck=True)
    assert result.created == 0 and result.price_changed == 1 and result.marked_unavailable == 1
    assert [p for _, p in storage.price_history(first[0].listing_id)] == [15000, 14000]
    rows = {r["listing_id"]: r for r in storage.listing_rows(statuses=[s.value for s in ListingStatus])}
    assert rows[first[2].listing_id]["status"] == ListingStatus.UNAVAILABLE.value
    # история анализов не перезаписывается
    assert len(storage.events(kind="campaign")) == 2
