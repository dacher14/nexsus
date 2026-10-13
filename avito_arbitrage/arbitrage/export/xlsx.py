"""Запись рабочего документа в XLSX (импортируется в Google Sheets без потерь формул)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .workbook import Sheet, WorkbookInput, build_sheets

NUMBER_FORMATS = {"rub": '#,##0" ₽";[Red]-#,##0" ₽"', "pct": "0%", "int": "0"}
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=12)
BANNER_FILL = PatternFill("solid", fgColor="FDE68A")
RED = PatternFill("solid", fgColor="FCA5A5")
AMBER = PatternFill("solid", fgColor="FDE68A")
GREEN = PatternFill("solid", fgColor="86EFAC")
LIGHT_GREEN = PatternFill("solid", fgColor="D9F99D")
VALIDATION_ROWS = 2000


def _write_sheet(ws, sheet: Sheet) -> None:
    row = 1
    if sheet.banner:
        ws.cell(row=1, column=1, value=sheet.banner).font = Font(bold=True)
        ws.cell(row=1, column=1).fill = BANNER_FILL
        row = 2
    header_row = row
    for t_index, table in enumerate(sheet.tables):
        if table.title:
            ws.cell(row=row, column=1, value=table.title).font = TITLE_FONT
            row += 1
        header_row = row
        for c, col in enumerate(table.columns, start=1):
            cell = ws.cell(row=row, column=c, value=col.header)
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            letter = get_column_letter(c)
            ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, col.width)
        row += 1
        for values in table.rows:
            for c, (col, value) in enumerate(zip(table.columns, values), start=1):
                cell = ws.cell(row=row, column=c, value=value)
                if col.fmt in NUMBER_FORMATS:
                    cell.number_format = NUMBER_FORMATS[col.fmt]
                elif col.fmt == "wrap":
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                elif col.fmt == "url" and isinstance(value, str) and value.startswith("http"):
                    cell.hyperlink = value
                    cell.font = Font(color="2563EB", underline="single")
            row += 1
        if t_index < len(sheet.tables) - 1:
            row += 1

    if len(sheet.tables) != 1:
        return
    table = sheet.tables[0]
    last_row = max(header_row + 1, header_row + len(table.rows))
    last_col = get_column_letter(len(table.columns))
    if sheet.freeze_header:
        ws.freeze_panes = ws.cell(row=header_row + 1, column=2 if len(table.columns) > 20 else 1)
    if sheet.autofilter:
        ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
    headers = [c.header for c in table.columns]
    for header, options in sheet.validations.items():
        letter = get_column_letter(headers.index(header) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.add(f"{letter}{header_row + 1}:{letter}{header_row + VALIDATION_ROWS}")
        ws.add_data_validation(dv)
    for header, rule in sheet.highlight.items():
        letter = get_column_letter(headers.index(header) + 1)
        rng = f"{letter}{header_row + 1}:{letter}{last_row}"
        if rule == "risk":
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["70"], fill=RED))
            ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["45", "69"], fill=AMBER))
        elif rule == "rating":
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["85"], fill=GREEN))
            ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["70", "84"], fill=LIGHT_GREEN))
        elif rule == "reco":
            first = f"{letter}{header_row + 1}"
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},11)="НЕ ПОКУПАТЬ"'], fill=RED))


def write_xlsx(data: WorkbookInput, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in build_sheets(data):
        _write_sheet(wb.create_sheet(sheet.name), sheet)
    wb.save(path)
    return path
