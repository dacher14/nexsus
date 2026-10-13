"""Выгрузка в Google Sheets через сервисный аккаунт (gspread).

Требуется: pip install gspread и JSON-ключ сервисного аккаунта с доступом к
Google Sheets/Drive API (GOOGLE_SERVICE_ACCOUNT_JSON). Таблица создаётся от
имени сервисного аккаунта и открывается пользователю по e-mail (share_with).
Если интеграция не настроена — используйте XLSX и импорт в Google Sheets.
"""
from __future__ import annotations

import os
from typing import Any

from openpyxl.utils import get_column_letter

from .workbook import Sheet, WorkbookInput, build_sheets


class SheetsNotConfigured(RuntimeError):
    pass


def _grid(sheet: Sheet) -> list[list[Any]]:
    grid: list[list[Any]] = []
    if sheet.banner:
        grid.append([sheet.banner])
    for i, table in enumerate(sheet.tables):
        if table.title:
            grid.append([table.title])
        grid.append([c.header for c in table.columns])
        grid.extend([["" if v is None else v for v in row] for row in table.rows])
        if i < len(sheet.tables) - 1:
            grid.append([])
    width = max((len(r) for r in grid), default=1)
    return [r + [""] * (width - len(r)) for r in grid]


def write_google_sheet(data: WorkbookInput, title: str, share_with: str | None = None,
                       credentials_path: str | None = None, spreadsheet_id: str | None = None) -> str:
    """Создаёт (или перезаписывает листы существующей) таблицы. Возвращает URL."""
    try:
        import gspread
    except ImportError as exc:
        raise SheetsNotConfigured("Не установлен gspread: pip install gspread") from exc
    credentials_path = credentials_path or os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not credentials_path or not os.path.exists(credentials_path):
        raise SheetsNotConfigured("Не задан ключ сервисного аккаунта (GOOGLE_SERVICE_ACCOUNT_JSON)")

    gc = gspread.service_account(filename=credentials_path)
    sh = gc.open_by_key(spreadsheet_id) if spreadsheet_id else gc.create(title)
    # Формулы записаны в синтаксисе en_US (запятые) — фиксируем локаль таблицы.
    sh.batch_update({"requests": [{"updateSpreadsheetProperties": {
        "properties": {"locale": "en_US"}, "fields": "locale"}}]})

    existing = {ws.title: ws for ws in sh.worksheets()}
    sheets = build_sheets(data)
    for sheet in sheets:
        grid = _grid(sheet)
        rows, cols = len(grid) + 50, max(len(grid[0]), 1)
        ws = existing.get(sheet.name)
        if ws is None:
            ws = sh.add_worksheet(title=sheet.name, rows=rows, cols=cols)
        else:
            ws.clear()
            ws.resize(rows=rows, cols=cols)
        ws.update(values=grid, range_name="A1", value_input_option="USER_ENTERED")
        if len(sheet.tables) == 1:
            header_row = 2 if sheet.banner else 1
            if sheet.freeze_header:
                ws.freeze(rows=header_row)
            if sheet.autofilter:
                ws.set_basic_filter(f"A{header_row}:{get_column_letter(cols)}{len(grid)}")
    ours = {s.name for s in sheets}
    for name, ws in existing.items():
        if name not in ours and name in ("Sheet1", "Лист1"):  # пустой лист по умолчанию
            sh.del_worksheet(ws)
    if share_with:
        sh.share(share_with, perm_type="user", role="writer", notify=True)
    return sh.url
