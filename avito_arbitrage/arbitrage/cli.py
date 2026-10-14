"""Командная строка: python -m arbitrage <команда> …"""
from __future__ import annotations

import argparse
import logging
import sys

from openpyxl import load_workbook

from .config import load_settings
from .export.workbook import USER_STATUSES
from .learning import update_deal
from .models import UserStatus
from .negotiation import Negotiator, NotAuthorized
from .pipeline import export_workbook, run_campaign, workbook_input
from .report import campaign_report, negotiation_report
from .search_plan import build_plan, write_plan
from .storage import Storage


def _storage(args, settings) -> Storage:
    return Storage(args.db or settings.database_url)


def cmd_campaign(args, settings) -> int:
    storage = _storage(args, settings)
    try:
        result = run_campaign(storage, settings, args.input, comps_path=args.comps, out_path=args.out,
                              use_llm=args.llm, full_recheck=args.full_recheck)
    except (FileNotFoundError, ValueError) as exc:
        print(f"Ошибка входных данных: {exc}", file=sys.stderr)
        return 2
    for err in result.errors[:20]:
        print(f"  пропущено: {err}", file=sys.stderr)
    data = workbook_input(storage, settings)
    if args.google_sheet:
        _push_sheets(data, args)
    print(campaign_report(data.analyses, result.stats(), str(result.workbook) if result.workbook else None,
                          settings, data.has_synthetic, statuses=data.statuses))
    return 0


def _push_sheets(data, args) -> None:
    from .export.gsheets import SheetsNotConfigured, write_google_sheet
    try:
        url = write_google_sheet(data, args.title, share_with=args.share, spreadsheet_id=args.spreadsheet_id)
        print(f"Google-таблица: {url}")
    except SheetsNotConfigured as exc:
        print(f"Google Sheets не подключён ({exc}); используйте XLSX → Файл → Импорт в Google Sheets.")


def cmd_export(args, settings) -> int:
    storage = _storage(args, settings)
    path = export_workbook(storage, settings, args.out)
    print(f"Таблица сохранена: {path}")
    if args.google_sheet:
        _push_sheets(workbook_input(storage, settings), args)
    return 0


def cmd_status(args, settings) -> int:
    storage = _storage(args, settings)
    for lid in args.ids:
        storage.set_user_status(lid, UserStatus(args.status), args.note or "")
        print(f"{lid}: {args.status}")
    return 0


def cmd_import_statuses(args, settings) -> int:
    storage = _storage(args, settings)
    ws = load_workbook(args.xlsx, read_only=True)["ALL DEALS"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    id_col, st_col = header.index("ID записи"), header.index("Статус пользователя")
    current = storage.user_statuses()
    changed = 0
    for row in rows:
        lid, status = row[id_col], str(row[st_col]).strip() if row[st_col] is not None else ""
        if not lid or not status:
            continue
        if status not in USER_STATUSES:
            print(f"  {lid}: неизвестный статус «{status}» — пропущено", file=sys.stderr)
            continue
        if current.get(str(lid)) != status:
            storage.set_user_status(str(lid), UserStatus(status), "импорт из таблицы")
            changed += 1
    print(f"Обновлено статусов: {changed}. Одобрено для переговоров: "
          f"{sum(1 for s in storage.user_statuses().values() if s == UserStatus.APPROVED.value)}")
    print("Новые одобренные объявления попадут в переговоры только после отдельной явной команды (negotiate authorize).")
    return 0


def cmd_negotiate(args, settings) -> int:
    storage = _storage(args, settings)
    neg = Negotiator(storage, settings)
    try:
        if args.action == "authorize":
            ids = neg.authorize(args.text)
            print(f"Разрешение получено для {len(ids)} одобренных объявлений: {', '.join(ids) or '—'}")
        elif args.action == "prepare":
            ids = args.ids or [lid for lid, s in storage.negotiations().items() if s.get("authorized_at")]
            if not ids:
                print("Нет разрешённых переговоров: одобрите объявления (APPROVED) и дайте команду authorize.")
            for lid in ids:
                _, draft, note = neg.prepare(lid)
                print(f"\n=== {lid}: {note}")
                if draft:
                    print(draft)
        elif args.action == "sent":
            for lid in args.ids:
                print(f"{lid}: {'отправка подтверждена' if neg.mark_sent(lid) else 'нет неотправленного черновика'}")
        elif args.action == "reply":
            out = neg.process_reply(args.ids[0], args.text)
            print(f"Позиция продавца: {out.position}; статус: {out.status.value}\nСледующий шаг: {out.next_step}")
            for fact in out.new_facts:
                print(f"  новые сведения: {fact}")
            if out.draft:
                print(f"\nЧерновик ответа (отправьте вручную):\n{out.draft}")
    except (NotAuthorized, KeyError) as exc:
        print(f"Отказано: {exc}", file=sys.stderr)
        return 2
    if args.out:
        export_workbook(storage, settings, args.out)
    return 0


def cmd_deal(args, settings) -> int:
    storage = _storage(args, settings)
    if storage.get_listing(args.id) is None:
        print(f"Объявление {args.id} не найдено в базе", file=sys.stderr)
        return 2
    data = update_deal(storage, args.id, args.status, purchase_price=args.purchase, costs_so_far=args.costs,
                       repair_cost=args.repair_cost, repair_hours=args.repair_hours, actual_fault=args.fault,
                       bought_at=args.bought_at, listed_at=args.listed_at, sold_price=args.sold,
                       sold_at=args.sold_at, notes=args.notes)
    print(f"{args.id}: {args.status}")
    if "net_profit" in data:
        print(f"Фактическая чистая прибыль: {data['net_profit']} ₽; ошибки прогноза: {data.get('errors', {})}")
    return 0


def cmd_report(args, settings) -> int:
    storage = _storage(args, settings)
    data = workbook_input(storage, settings)
    print(campaign_report(data.analyses, data.campaign, None, settings, data.has_synthetic,
                          statuses=data.statuses))
    print()
    print(negotiation_report(storage))
    return 0


def cmd_search_plan(args, settings) -> int:
    cities = args.cities if args.cities is not None else (settings.search_cities or None)
    plan = build_plan(args.categories, [("" if c.lower() in ("россия", "вся россия") else c) for c in cities]
                      if cities else None, home=settings.region)
    path = write_plan(plan, args.out)
    print(f"План поиска: {len(plan)} запросов → {path}")
    return 0


def cmd_telegram(args, settings) -> int:
    from .telegram_bot import run_bot
    run_bot(_storage(args, settings), settings, args.out)
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="arbitrage", description="AVITO ARBITRAGE AI — поиск и оценка сделок")
    p.add_argument("--config", help="YAML с настройками (см. config.example.yaml)")
    p.add_argument("--db", help="URL базы данных (по умолчанию из настроек / DATABASE_URL)")
    p.add_argument("-v", "--verbose", action="store_true")
    sub = p.add_subparsers(dest="command", required=True)

    def sheets_opts(sp):
        sp.add_argument("--google-sheet", action="store_true", help="также выгрузить в Google Sheets")
        sp.add_argument("--share", help="e-mail, которому открыть Google-таблицу")
        sp.add_argument("--spreadsheet-id", help="перезаписать листы существующей Google-таблицы")
        sp.add_argument("--title", default="AVITO ARBITRAGE — кампания")

    c = sub.add_parser("campaign", help="исследовательская кампания по выгрузке объявлений")
    c.add_argument("--input", nargs="+", required=True, help="JSON/CSV с объявлениями")
    c.add_argument("--comps", help="JSON/CSV с дополнительными аналогами")
    c.add_argument("--out", default="out/avito_arbitrage.xlsx")
    c.add_argument("--llm", action="store_true", help="LLM-анализ неоднозначных объявлений (llm_enabled)")
    c.add_argument("--full-recheck", action="store_true",
                   help="отсутствующие в выгрузке объявления пометить недоступными")
    sheets_opts(c)
    c.set_defaults(func=cmd_campaign)

    e = sub.add_parser("export", help="пересобрать таблицу из базы")
    e.add_argument("--out", default="out/avito_arbitrage.xlsx")
    sheets_opts(e)
    e.set_defaults(func=cmd_export)

    s = sub.add_parser("status", help="статус пользователя для объявлений")
    s.add_argument("status", choices=USER_STATUSES)
    s.add_argument("ids", nargs="+")
    s.add_argument("--note")
    s.set_defaults(func=cmd_status)

    i = sub.add_parser("import-statuses", help="перенести статусы из отредактированной таблицы в базу")
    i.add_argument("xlsx")
    i.set_defaults(func=cmd_import_statuses)

    n = sub.add_parser("negotiate", help="переговоры (только после одобрения и явной команды)")
    n.add_argument("action", choices=["authorize", "prepare", "sent", "reply"])
    n.add_argument("ids", nargs="*", help="ID объявлений")
    n.add_argument("--text", help="фраза-разрешение (authorize) или ответ продавца (reply)")
    n.add_argument("--out", help="обновить таблицу после действия")
    n.set_defaults(func=cmd_negotiate)

    d = sub.add_parser("deal", help="фактические результаты сделки")
    d.add_argument("id")
    d.add_argument("--status", required=True, choices=["bought", "repairing", "listed", "sold", "written_off"])
    d.add_argument("--purchase", type=int)
    d.add_argument("--costs", type=int, help="все расходы кроме покупки")
    d.add_argument("--repair-cost", type=int)
    d.add_argument("--repair-hours", type=float)
    d.add_argument("--fault", help="фактическая неисправность")
    d.add_argument("--bought-at")
    d.add_argument("--listed-at")
    d.add_argument("--sold", type=int)
    d.add_argument("--sold-at")
    d.add_argument("--notes")
    d.set_defaults(func=cmd_deal)

    r = sub.add_parser("report", help="итоговый отчёт")
    r.set_defaults(func=cmd_report)

    sp = sub.add_parser("search-plan", help="план многопроходного поиска (CSV)")
    sp.add_argument("--out", default="out/search_plan.csv")
    sp.add_argument("--categories", nargs="*")
    sp.add_argument("--cities", nargs="*", help="города поиска; «Россия» — все регионы с доставкой")
    sp.set_defaults(func=cmd_search_plan)

    t = sub.add_parser("telegram", help="Telegram-бот управления (TELEGRAM_BOT_TOKEN, TELEGRAM_ALLOWED_CHAT_ID)")
    t.add_argument("--out", default="out/avito_arbitrage.xlsx")
    t.set_defaults(func=cmd_telegram)
    return p


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    logging.basicConfig(level=logging.INFO if args.verbose else logging.WARNING, format="%(levelname)s %(message)s")
    settings = load_settings(args.config)
    if args.command == "negotiate" and args.action in ("authorize", "reply") and not args.text:
        print("Нужен параметр --text", file=sys.stderr)
        return 2
    if args.command == "negotiate" and args.action in ("sent", "reply") and not args.ids:
        print("Укажите ID объявления", file=sys.stderr)
        return 2
    return args.func(args, settings)
